from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


CHANNELS = ("境内新闻", "公众文章", "今日头条", "微博", "小红书", "抖音")
REQUIRED_ANY = {"标题", "正文内容", "视频ASR"}


def serial(value: object) -> object:
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    return value


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def infer_channel(path: Path, sheet_name: str) -> str:
    joined = f"{path.name} {sheet_name}"
    for channel in CHANNELS:
        if channel in joined:
            return channel
    if "公号" in joined or "微信" in joined:
        return "公众文章"
    if "新闻" in joined:
        return "境内新闻"
    raise ValueError(f"无法识别渠道，请修正文件名或工作表名：{path} / {sheet_name}")


def stable_id(relative: str, sheet: str, row_number: int, data: dict) -> str:
    native = str(data.get("发文ID") or data.get("内容ID") or "").strip()
    seed = f"{relative}|{sheet}|{row_number}|{native}"
    return hashlib.sha256(seed.encode("utf-8")).hexdigest()[:16]


def extract(source: Path, output: Path) -> dict:
    source = source.resolve()
    output = output.resolve()
    if not source.is_dir():
        raise FileNotFoundError(source)
    if output == source or source in output.parents:
        raise ValueError("输出目录不能位于原始数据目录内")
    output.mkdir(parents=True, exist_ok=True)

    records: list[dict] = []
    manifest: list[dict] = []
    for path in sorted(source.rglob("*")):
        if not path.is_file() or path.name.startswith("~$"):
            continue
        relative = str(path.relative_to(source))
        item = {
            "path": str(path),
            "relative": relative,
            "batch": path.parent.name,
            "bytes": path.stat().st_size,
            "sha256": file_sha256(path),
            "used_for_screening": path.suffix.lower() in {".xlsx", ".xlsm"},
        }
        manifest.append(item)
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = []
        try:
            for worksheet in workbook.worksheets:
                declared = worksheet.calculate_dimension()
                worksheet.reset_dimensions()
                rows = [
                    (index, tuple(serial(value) for value in row))
                    for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1)
                    if any(value is not None for value in row)
                ]
                if not rows:
                    sheets.append({"name": worksheet.title, "rows": 0, "declared_dimension": declared})
                    continue
                header_row, raw_headers = rows[0]
                headers = [str(value).strip() if value is not None else f"未命名列{index + 1}" for index, value in enumerate(raw_headers)]
                if not REQUIRED_ANY.intersection(headers):
                    raise ValueError(f"表头缺少标题、正文内容或视频ASR：{path.name} / {worksheet.title}")
                channel = infer_channel(path, worksheet.title)
                count = 0
                for row_number, raw_row in rows[1:]:
                    data = {headers[index]: value for index, value in enumerate(raw_row) if index < len(headers)}
                    record_id = stable_id(relative, worksheet.title, row_number, data)
                    records.append({
                        "id": record_id,
                        "batch": path.parent.name,
                        "channel": channel,
                        "source_file": path.name,
                        "source_path": str(path),
                        "sheet": worksheet.title,
                        "row": row_number,
                        "data": data,
                    })
                    count += 1
                sheets.append({"name": worksheet.title, "rows": count, "header_row": header_row, "declared_dimension": declared})
        finally:
            workbook.close()
        item["sheets"] = sheets

    if not records:
        raise RuntimeError("没有从Excel中读取到任何样本")
    ids = [record["id"] for record in records]
    if len(ids) != len(set(ids)):
        raise RuntimeError("稳定来源ID发生重复")

    records_path = output / "records.jsonl"
    records_path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in records), encoding="utf-8")
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    summary = {
        "status": "PASS",
        "source": str(source),
        "records": len(records),
        "batches": sorted({row["batch"] for row in records}),
        "channels": {channel: sum(row["channel"] == channel for row in records) for channel in CHANNELS},
        "reports_or_other_files_used_for_screening": 0,
        "records_path": str(records_path),
    }
    (output / "extraction_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract film-brief source rows from channel workbooks")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    extract(args.source, args.output)


if __name__ == "__main__":
    main()
